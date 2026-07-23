"""
Export routes - Excel export functionality.
"""
from flask import Blueprint, render_template, request, session, send_file, flash, redirect, url_for
from datetime import datetime
from io import BytesIO
from utils.auth import login_required, get_current_user
from utils.jalali import format_jalali_date, format_jalali_datetime, get_now_jalali
from models import db, User, Project, DailyFormSubmission
from translations import get_translation
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

export_bp = Blueprint('export', __name__)


@export_bp.route('/')
@login_required
def export_page():
    """Export options page."""
    user = get_current_user()
    lang = session.get('lang', 'en')
    
    projects = Project.query.filter_by(is_active=True).all()
    
    # Employees only export their own data
    employees = []
    if user['role'] in ['admin', 'boss']:
        employees = User.query.filter_by(role='employee').all()
    
    return render_template('export/index.html', lang=lang, projects=projects, employees=employees)


@export_bp.route('/forms')
@login_required
def export_forms():
    """Export forms to Excel."""
    user = get_current_user()
    
    # Get filters
    from_date = request.args.get('from_date')
    to_date = request.args.get('to_date')
    project_id = request.args.get('project_id', type=int)
    employee_id = request.args.get('employee_id', type=int)
    status_filter = request.args.get('status', '')
    
    db_user = User.query.filter_by(username=user['username']).first()
    
    # Build query
    query = DailyFormSubmission.query
    
    # Role-based filtering
    if user['role'] == 'employee':
        query = query.filter_by(submitted_by=db_user.id)
    elif employee_id:
        query = query.filter_by(submitted_by=employee_id)
    
    # Apply filters
    if from_date:
        query = query.filter(DailyFormSubmission.form_date >= from_date)
    if to_date:
        query = query.filter(DailyFormSubmission.form_date <= to_date)
    if project_id:
        query = query.filter_by(project_id=project_id)
    if status_filter:
        query = query.filter_by(status=status_filter)
    
    forms = query.order_by(DailyFormSubmission.form_date.desc()).all()
    
    if not forms:
        lang = session.get('lang', 'en')
        t = lambda key: get_translation(key, lang)
        flash(t('no_forms_found'), 'warning')
        return redirect(url_for('export.export_page'))
    
    # Create Excel workbook
    wb = create_forms_workbook(forms)
    
    # Save to BytesIO
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    # Use Jalali date in filename
    jalali_now = get_now_jalali()
    filename = f"daily_forms_export_{jalali_now.strftime('%Y%m%d_%H%M%S')}.xlsx"
    
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename
    )


@export_bp.route('/form/<int:form_id>')
@login_required
def export_single_form(form_id):
    """Export a single form to Excel."""
    user = get_current_user()
    
    form = DailyFormSubmission.query.get_or_404(form_id)
    
    # Check permission
    db_user = User.query.filter_by(username=user['username']).first()
    if user['role'] == 'employee' and form.submitted_by != db_user.id:
        lang = session.get('lang', 'en')
        t = lambda key: get_translation(key, lang)
        flash(t('access_denied'), 'danger')
        return redirect(url_for('main.dashboard'))
    
    # Create Excel workbook
    wb = create_single_form_workbook(form)
    
    # Save to BytesIO
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    # Use Jalali date in filename
    jalali_date = format_jalali_date(form.form_date, '%Y%m%d')
    filename = f"form_{form.document_code}_{jalali_date}.xlsx"
    
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename
    )


def create_forms_workbook(forms):
    """Create Excel workbook with multiple forms."""
    wb = openpyxl.Workbook()
    
    # Styles
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')
    header_alignment = Alignment(horizontal='center', vertical='center')
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Summary sheet
    ws = wb.active
    ws.title = "Summary"
    
    headers = ['Document Code', 'Date', 'Project', 'Submitted By', 'Status', 'Submitted At']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    
    for row, form in enumerate(forms, 2):
        ws.cell(row=row, column=1, value=form.document_code).border = thin_border
        ws.cell(row=row, column=2, value=format_jalali_date(form.form_date)).border = thin_border
        ws.cell(row=row, column=3, value=form.project.name if form.project else '').border = thin_border
        ws.cell(row=row, column=4, value=form.submitted_by_user.username if form.submitted_by_user else '').border = thin_border
        ws.cell(row=row, column=5, value=form.status).border = thin_border
        ws.cell(row=row, column=6, value=format_jalali_datetime(form.submitted_at) if form.submitted_at else '').border = thin_border
    
    # Adjust column widths
    for col in range(1, 7):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 20
    
    # Human Resources sheet
    ws_hr = wb.create_sheet("Human Resources")
    hr_headers = ['Document Code', 'Date', 'Position', 'Count', 'Working Hours', 'Notes']
    for col, header in enumerate(hr_headers, 1):
        cell = ws_hr.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    
    hr_row = 2
    for form in forms:
        for hr in form.human_resources:
            ws_hr.cell(row=hr_row, column=1, value=form.document_code).border = thin_border
            ws_hr.cell(row=hr_row, column=2, value=format_jalali_date(form.form_date)).border = thin_border
            ws_hr.cell(row=hr_row, column=3, value=hr.post).border = thin_border
            ws_hr.cell(row=hr_row, column=4, value=hr.count).border = thin_border
            ws_hr.cell(row=hr_row, column=5, value=hr.working_hours).border = thin_border
            ws_hr.cell(row=hr_row, column=6, value=hr.notes or '').border = thin_border
            hr_row += 1
    
    # Equipment sheet
    ws_eq = wb.create_sheet("Equipment")
    eq_headers = ['Document Code', 'Date', 'Equipment', 'Active Count', 'Working Hours', 'Situation', 'Reason']
    for col, header in enumerate(eq_headers, 1):
        cell = ws_eq.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    
    eq_row = 2
    for form in forms:
        for eq in form.tools_equipments:
            ws_eq.cell(row=eq_row, column=1, value=form.document_code).border = thin_border
            ws_eq.cell(row=eq_row, column=2, value=format_jalali_date(form.form_date)).border = thin_border
            ws_eq.cell(row=eq_row, column=3, value=eq.equipment_name).border = thin_border
            ws_eq.cell(row=eq_row, column=4, value=eq.count_active).border = thin_border
            ws_eq.cell(row=eq_row, column=5, value=eq.working_hours).border = thin_border
            ws_eq.cell(row=eq_row, column=6, value=eq.situation or '').border = thin_border
            ws_eq.cell(row=eq_row, column=7, value=eq.inactivity_reason or '').border = thin_border
            eq_row += 1
    
    # Operations sheet
    ws_op = wb.create_sheet("Operations")
    op_headers = ['Document Code', 'Date', 'Operation', 'Start KM', 'End KM', 'Unit', 'Progress', 'Map Number']
    for col, header in enumerate(op_headers, 1):
        cell = ws_op.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    
    op_row = 2
    for form in forms:
        for op in form.construction_operations:
            ws_op.cell(row=op_row, column=1, value=form.document_code).border = thin_border
            ws_op.cell(row=op_row, column=2, value=format_jalali_date(form.form_date)).border = thin_border
            ws_op.cell(row=op_row, column=3, value=op.operation_type).border = thin_border
            ws_op.cell(row=op_row, column=4, value=op.start_km).border = thin_border
            ws_op.cell(row=op_row, column=5, value=op.end_km).border = thin_border
            ws_op.cell(row=op_row, column=6, value=op.unit or '').border = thin_border
            ws_op.cell(row=op_row, column=7, value=op.progress).border = thin_border
            ws_op.cell(row=op_row, column=8, value=op.map_number or '').border = thin_border
            op_row += 1
    
    # Safety sheet
    ws_safety = wb.create_sheet("Safety")
    safety_headers = ['Document Code', 'Date', 'Safety Situation', 'Inspection Done', 'Incident Occurred', 'Explanation']
    for col, header in enumerate(safety_headers, 1):
        cell = ws_safety.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    
    safety_row = 2
    for form in forms:
        for safety in form.safety_records:
            ws_safety.cell(row=safety_row, column=1, value=form.document_code).border = thin_border
            ws_safety.cell(row=safety_row, column=2, value=format_jalali_date(form.form_date)).border = thin_border
            ws_safety.cell(row=safety_row, column=3, value=safety.safety_situation or '').border = thin_border
            ws_safety.cell(row=safety_row, column=4, value='Yes' if safety.safety_inspection else 'No').border = thin_border
            ws_safety.cell(row=safety_row, column=5, value='Yes' if safety.incident_occurred else 'No').border = thin_border
            ws_safety.cell(row=safety_row, column=6, value=safety.incident_explanation or '').border = thin_border
            safety_row += 1
    
    return wb


def create_single_form_workbook(form):
    """Create Excel workbook for a single form."""
    wb = openpyxl.Workbook()
    
    # Styles
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')
    section_fill = PatternFill(start_color='E2E8F0', end_color='E2E8F0', fill_type='solid')
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    ws = wb.active
    ws.title = "Daily Form"
    
    row = 1
    
    # Header info
    ws.cell(row=row, column=1, value="Document Code:").font = Font(bold=True)
    ws.cell(row=row, column=2, value=form.document_code)
    row += 1
    
    ws.cell(row=row, column=1, value="Date:").font = Font(bold=True)
    ws.cell(row=row, column=2, value=format_jalali_date(form.form_date))
    row += 1
    
    ws.cell(row=row, column=1, value="Project:").font = Font(bold=True)
    ws.cell(row=row, column=2, value=form.project.name if form.project else '')
    row += 1
    
    ws.cell(row=row, column=1, value="Day of Week:").font = Font(bold=True)
    ws.cell(row=row, column=2, value=form.day_of_week or '')
    row += 1
    
    ws.cell(row=row, column=1, value="Work Shift:").font = Font(bold=True)
    ws.cell(row=row, column=2, value=form.work_shift or '')
    row += 2
    
    # Human Resources section
    ws.cell(row=row, column=1, value="HUMAN RESOURCES").font = Font(bold=True, size=12)
    ws.cell(row=row, column=1).fill = section_fill
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    row += 1
    
    hr_headers = ['Position', 'Count', 'Working Hours', 'Notes']
    for col, header in enumerate(hr_headers, 1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border
    row += 1
    
    for hr in form.human_resources:
        ws.cell(row=row, column=1, value=hr.post).border = thin_border
        ws.cell(row=row, column=2, value=hr.count).border = thin_border
        ws.cell(row=row, column=3, value=hr.working_hours).border = thin_border
        ws.cell(row=row, column=4, value=hr.notes or '').border = thin_border
        row += 1
    
    row += 1
    
    # Equipment section
    ws.cell(row=row, column=1, value="TOOLS AND EQUIPMENT").font = Font(bold=True, size=12)
    ws.cell(row=row, column=1).fill = section_fill
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    row += 1
    
    eq_headers = ['Equipment', 'Active', 'Hours', 'Situation', 'Reason']
    for col, header in enumerate(eq_headers, 1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border
    row += 1
    
    for eq in form.tools_equipments:
        ws.cell(row=row, column=1, value=eq.equipment_name).border = thin_border
        ws.cell(row=row, column=2, value=eq.count_active).border = thin_border
        ws.cell(row=row, column=3, value=eq.working_hours).border = thin_border
        ws.cell(row=row, column=4, value=eq.situation or '').border = thin_border
        ws.cell(row=row, column=5, value=eq.inactivity_reason or '').border = thin_border
        row += 1
    
    # Adjust column widths
    ws.column_dimensions['A'].width = 25
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 30
    
    return wb
