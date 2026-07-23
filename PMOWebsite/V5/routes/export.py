"""
Export routes - Excel export functionality.
"""
from flask import Blueprint, render_template, request, session, send_file, flash, redirect, url_for
from datetime import datetime
from io import BytesIO
from utils.auth import login_required, get_current_user
from utils.jalali import format_jalali_date, format_jalali_datetime, get_now_jalali
from models import db, User, Project, DailyFormSubmission
from translations import get_translation
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

export_bp = Blueprint('export', __name__)


@export_bp.route('/')
@login_required
def export_page():
    """Export options page."""
    user = get_current_user()
    lang = session.get('lang', 'en')
    
    projects = Project.query.filter_by(status='active').all()
    
    # Employees only export their own data
    employees = []
    if user['role'] in ['admin', 'boss']:
        employees = User.query.filter_by(role='employee').all()
    
    return render_template('export/index.html', lang=lang, projects=projects, employees=employees)


@export_bp.route('/forms')
@login_required
def export_forms():
    """Export forms to Excel."""
    user = get_current_user()
    
    # Get filters
    from_date = request.args.get('from_date')
    to_date = request.args.get('to_date')
    project_id = request.args.get('project_id', type=int)
    employee_id = request.args.get('employee_id', type=int)
    status_filter = request.args.get('status', '')
    
    db_user = User.query.filter_by(username=user['username']).first()
    
    # Build query
    query = DailyFormSubmission.query
    
    # Role-based filtering
    if user['role'] == 'employee':
        query = query.filter_by(submitted_by=db_user.id)
    elif employee_id:
        query = query.filter_by(submitted_by=employee_id)
    
    # Apply filters
    if from_date:
        query = query.filter(DailyFormSubmission.form_date >= from_date)
    if to_date:
        query = query.filter(DailyFormSubmission.form_date <= to_date)
    if project_id:
        query = query.filter_by(project_id=project_id)
    if status_filter:
        query = query.filter_by(status=status_filter)
    
    forms = query.order_by(DailyFormSubmission.form_date.desc()).all()
    
    if not forms:
        lang = session.get('lang', 'en')
        t = lambda key: get_translation(key, lang)
        flash(t('no_forms_found'), 'warning')
        return redirect(url_for('export.export_page'))
    
    # Create Excel workbook
    wb = create_forms_workbook(forms)
    
    # Save to BytesIO
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    # Use Jalali date in filename
    jalali_now = get_now_jalali()
    filename = f"daily_forms_export_{jalali_now.strftime('%Y%m%d_%H%M%S')}.xlsx"
    
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename
    )


@export_bp.route('/form/<int:form_id>')
@login_required
def export_single_form(form_id):
    """Export a single form to Excel."""
    user = get_current_user()
    
    form = DailyFormSubmission.query.get_or_404(form_id)
    
    # Check permission
    db_user = User.query.filter_by(username=user['username']).first()
    if user['role'] == 'employee' and form.submitted_by != db_user.id:
        lang = session.get('lang', 'en')
        t = lambda key: get_translation(key, lang)
        flash(t('access_denied'), 'danger')
        return redirect(url_for('main.dashboard'))
    
    # Create Excel workbook
    wb = create_single_form_workbook(form)
    
    # Save to BytesIO
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    # Use document code as filename (e.g., Project-03-keivan-14040801-001.xlsx)
    filename = f"{form.document_code}.xlsx"
    
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename
    )


def create_forms_workbook(forms):
    """Create Excel workbook with multiple forms."""
    wb = openpyxl.Workbook()
    
    # Styles
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')
    header_alignment = Alignment(horizontal='center', vertical='center')
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Summary sheet
    ws = wb.active
    ws.title = "Summary"
    
    headers = ['Document Code', 'Date', 'Project', 'Submitted By', 'Status', 'Submitted At']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    
    for row, form in enumerate(forms, 2):
        ws.cell(row=row, column=1, value=form.document_code).border = thin_border
        ws.cell(row=row, column=2, value=format_jalali_date(form.form_date)).border = thin_border
        ws.cell(row=row, column=3, value=form.project.name if form.project else '').border = thin_border
        ws.cell(row=row, column=4, value=form.submitted_by_user.username if form.submitted_by_user else '').border = thin_border
        ws.cell(row=row, column=5, value=form.status).border = thin_border
        ws.cell(row=row, column=6, value=format_jalali_datetime(form.submitted_at) if form.submitted_at else '').border = thin_border
    
    # Adjust column widths
    for col in range(1, 7):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 20
    
    # Human Resources sheet
    ws_hr = wb.create_sheet("Human Resources")
    hr_headers = ['Document Code', 'Date', 'Position', 'Count', 'Working Hours', 'Notes']
    for col, header in enumerate(hr_headers, 1):
        cell = ws_hr.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    
    hr_row = 2
    for form in forms:
        for hr in form.human_resources:
            ws_hr.cell(row=hr_row, column=1, value=form.document_code).border = thin_border
            ws_hr.cell(row=hr_row, column=2, value=format_jalali_date(form.form_date)).border = thin_border
            ws_hr.cell(row=hr_row, column=3, value=hr.post).border = thin_border
            ws_hr.cell(row=hr_row, column=4, value=hr.count).border = thin_border
            ws_hr.cell(row=hr_row, column=5, value=hr.working_hours).border = thin_border
            ws_hr.cell(row=hr_row, column=6, value=hr.notes or '').border = thin_border
            hr_row += 1
    
    # Equipment sheet
    ws_eq = wb.create_sheet("Equipment")
    eq_headers = ['Document Code', 'Date', 'Equipment', 'Active Count', 'Working Hours', 'Situation', 'Reason']
    for col, header in enumerate(eq_headers, 1):
        cell = ws_eq.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    
    eq_row = 2
    for form in forms:
        for eq in form.tools_equipments:
            ws_eq.cell(row=eq_row, column=1, value=form.document_code).border = thin_border
            ws_eq.cell(row=eq_row, column=2, value=format_jalali_date(form.form_date)).border = thin_border
            ws_eq.cell(row=eq_row, column=3, value=eq.equipment_name).border = thin_border
            ws_eq.cell(row=eq_row, column=4, value=eq.count_active).border = thin_border
            ws_eq.cell(row=eq_row, column=5, value=eq.working_hours).border = thin_border
            ws_eq.cell(row=eq_row, column=6, value=eq.situation or '').border = thin_border
            ws_eq.cell(row=eq_row, column=7, value=eq.inactivity_reason or '').border = thin_border
            eq_row += 1
    
    # Operations sheet
    ws_op = wb.create_sheet("Operations")
    op_headers = ['Document Code', 'Date', 'Operation', 'Start KM', 'End KM', 'Unit', 'Progress', 'Map Number']
    for col, header in enumerate(op_headers, 1):
        cell = ws_op.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    
    op_row = 2
    for form in forms:
        for op in form.construction_operations:
            ws_op.cell(row=op_row, column=1, value=form.document_code).border = thin_border
            ws_op.cell(row=op_row, column=2, value=format_jalali_date(form.form_date)).border = thin_border
            ws_op.cell(row=op_row, column=3, value=op.operation_type).border = thin_border
            ws_op.cell(row=op_row, column=4, value=op.start_station or '').border = thin_border
            ws_op.cell(row=op_row, column=5, value=op.end_station or '').border = thin_border
            ws_op.cell(row=op_row, column=6, value=op.unit or '').border = thin_border
            ws_op.cell(row=op_row, column=7, value=op.amount or 0).border = thin_border
            op_row += 1
    
    # Safety sheet
    ws_safety = wb.create_sheet("Safety")
    safety_headers = ['Document Code', 'Date', 'Safety Situation', 'Inspection Done', 'Incident Occurred', 'Explanation']
    for col, header in enumerate(safety_headers, 1):
        cell = ws_safety.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    
    safety_row = 2
    for form in forms:
        for safety in form.safety_records:
            ws_safety.cell(row=safety_row, column=1, value=form.document_code).border = thin_border
            ws_safety.cell(row=safety_row, column=2, value=format_jalali_date(form.form_date)).border = thin_border
            ws_safety.cell(row=safety_row, column=3, value=safety.safety_situation or '').border = thin_border
            ws_safety.cell(row=safety_row, column=4, value='Yes' if safety.safety_inspection else 'No').border = thin_border
            ws_safety.cell(row=safety_row, column=5, value='Yes' if safety.incident_occurred else 'No').border = thin_border
            ws_safety.cell(row=safety_row, column=6, value=safety.incident_explanation or '').border = thin_border
            safety_row += 1
    
    return wb


def create_single_form_workbook(form):
    """Create Excel workbook for a single form in the same format as upload template."""
    from openpyxl.utils import get_column_letter
    
    wb = openpyxl.Workbook()
    
    # Helper functions to convert English values to Persian for display
    def shift_to_persian(val):
        if not val:
            return ''
        val = val.lower()
        if val == 'morning':
            return 'صبح'
        elif val == 'night':
            return 'شب'
        return val
    
    def eq_situation_to_persian(val):
        if not val:
            return ''
        val_lower = val.lower()
        if val_lower == 'active':
            return 'فعال'
        elif val_lower == 'inactive':
            return 'غیرفعال'
        elif 'repair' in val_lower:
            return 'در حال تعمیر'
        return val
    
    def weather_to_persian(val):
        if not val:
            return ''
        val = val.lower()
        mapping = {
            'clear': 'صاف',
            'cloudy': 'ابری',
            'rainy': 'بارانی',
            'foggy': 'مه‌آلود',
            'snowy': 'برفی'
        }
        return mapping.get(val, val)
    
    def wind_to_persian(val):
        if not val:
            return ''
        val = val.lower()
        mapping = {
            'slow': 'آرام',
            'normal': 'معمولی',
            'fast': 'تند'
        }
        return mapping.get(val, val)
    
    def effect_to_persian(val):
        if not val:
            return ''
        val = val.lower()
        mapping = {
            'no_effect': 'بدون تاثیر',
            'slow_progress': 'کندی پیشرفت',
            'stopped': 'توقف کار'
        }
        return mapping.get(val, val)
    
    def safety_to_persian(val):
        if not val:
            return ''
        val_lower = val.lower()
        mapping = {
            'good': 'خوب',
            'fair': 'متوسط',
            'poor': 'ضعیف'
        }
        return mapping.get(val_lower, val)
    
    # Styles - same as upload template
    header_font = Font(bold=True, size=12, color="FFFFFF")
    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    section_font = Font(bold=True, size=14, color="FFFFFF")
    section_fill = PatternFill(start_color="059669", end_color="059669", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    right_align = Alignment(horizontal='right', vertical='center', wrap_text=True)
    
    ws = wb.active
    ws.title = "فرم روزانه"
    ws.sheet_view.rightToLeft = True  # RTL for Persian
    
    current_row = 1
    
    # ==================== Section 1: Date and Time ====================
    ws.merge_cells(f'A{current_row}:D{current_row}')
    ws[f'A{current_row}'] = "۱. تاریخ و زمان"
    ws[f'A{current_row}'].font = section_font
    ws[f'A{current_row}'].fill = section_fill
    ws[f'A{current_row}'].alignment = center_align
    current_row += 1
    
    # Date and Time headers
    headers = ["تاریخ (1404/01/01)", "روز هفته", "شیفت کاری", "پروژه"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=current_row, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col)].width = 20
    current_row += 1
    
    # Data row
    ws.cell(row=current_row, column=1, value=format_jalali_date(form.form_date)).border = thin_border
    ws.cell(row=current_row, column=1).alignment = center_align
    ws.cell(row=current_row, column=2, value=form.day_of_week or '').border = thin_border
    ws.cell(row=current_row, column=2).alignment = center_align
    ws.cell(row=current_row, column=3, value=shift_to_persian(form.work_shift)).border = thin_border
    ws.cell(row=current_row, column=3).alignment = center_align
    ws.cell(row=current_row, column=4, value=form.project.name if form.project else '').border = thin_border
    ws.cell(row=current_row, column=4).alignment = center_align
    current_row += 2
    
    # ==================== Section 2: Human Resources ====================
    ws.merge_cells(f'A{current_row}:D{current_row}')
    ws[f'A{current_row}'] = "۲. نیروی انسانی"
    ws[f'A{current_row}'].font = section_font
    ws[f'A{current_row}'].fill = section_fill
    ws[f'A{current_row}'].alignment = center_align
    current_row += 1
    
    # HR headers
    hr_headers = ["سمت", "تعداد", "ساعت کار", "توضیحات"]
    for col, header in enumerate(hr_headers, 1):
        cell = ws.cell(row=current_row, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    current_row += 1
    
    # HR data
    for hr in form.human_resources:
        ws.cell(row=current_row, column=1, value=hr.post).border = thin_border
        ws.cell(row=current_row, column=1).alignment = right_align
        ws.cell(row=current_row, column=2, value=hr.count).border = thin_border
        ws.cell(row=current_row, column=2).alignment = center_align
        ws.cell(row=current_row, column=3, value=hr.working_hours).border = thin_border
        ws.cell(row=current_row, column=3).alignment = center_align
        ws.cell(row=current_row, column=4, value=hr.notes or '').border = thin_border
        ws.cell(row=current_row, column=4).alignment = center_align
        current_row += 1
    
    current_row += 1
    
    # ==================== Section 3: Tools and Equipment ====================
    ws.merge_cells(f'A{current_row}:E{current_row}')
    ws[f'A{current_row}'] = "۳. ماشین‌آلات و تجهیزات"
    ws[f'A{current_row}'].font = section_font
    ws[f'A{current_row}'].fill = section_fill
    ws[f'A{current_row}'].alignment = center_align
    current_row += 1
    
    # Equipment headers
    eq_headers = ["نام تجهیزات", "تعداد", "ساعت کار", "وضعیت", "دلیل غیرفعال"]
    for col, header in enumerate(eq_headers, 1):
        cell = ws.cell(row=current_row, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    current_row += 1
    
    # Equipment data
    for eq in form.tools_equipments:
        ws.cell(row=current_row, column=1, value=eq.equipment_name).border = thin_border
        ws.cell(row=current_row, column=1).alignment = right_align
        ws.cell(row=current_row, column=2, value=eq.count_active).border = thin_border
        ws.cell(row=current_row, column=2).alignment = center_align
        ws.cell(row=current_row, column=3, value=eq.working_hours).border = thin_border
        ws.cell(row=current_row, column=3).alignment = center_align
        ws.cell(row=current_row, column=4, value=eq_situation_to_persian(eq.situation)).border = thin_border
        ws.cell(row=current_row, column=4).alignment = center_align
        ws.cell(row=current_row, column=5, value=eq.inactivity_reason or '').border = thin_border
        ws.cell(row=current_row, column=5).alignment = center_align
        current_row += 1
    
    current_row += 1
    
    # ==================== Section 4: Construction Operations ====================
    ws.merge_cells(f'A{current_row}:F{current_row}')
    ws[f'A{current_row}'] = "۴. عملیات ساختمانی"
    ws[f'A{current_row}'].font = section_font
    ws[f'A{current_row}'].fill = section_fill
    ws[f'A{current_row}'].alignment = center_align
    current_row += 1
    
    # Operations headers
    op_headers = ["نوع عملیات", "کیلومتر شروع", "کیلومتر پایان", "واحد", "پیشرفت", "شماره نقشه"]
    for col, header in enumerate(op_headers, 1):
        cell = ws.cell(row=current_row, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    current_row += 1
    
    # Operations data
    for op in form.construction_operations:
        ws.cell(row=current_row, column=1, value=op.operation_type).border = thin_border
        ws.cell(row=current_row, column=1).alignment = right_align
        ws.cell(row=current_row, column=2, value=op.start_station or '').border = thin_border
        ws.cell(row=current_row, column=2).alignment = center_align
        ws.cell(row=current_row, column=3, value=op.end_station or '').border = thin_border
        ws.cell(row=current_row, column=3).alignment = center_align
        ws.cell(row=current_row, column=4, value=op.unit or '').border = thin_border
        ws.cell(row=current_row, column=4).alignment = center_align
        ws.cell(row=current_row, column=5, value=op.amount or 0).border = thin_border
        ws.cell(row=current_row, column=5).alignment = center_align
        current_row += 1
    
    current_row += 1
    
    # ==================== Section 5: Incoming Materials ====================
    ws.merge_cells(f'A{current_row}:G{current_row}')
    ws[f'A{current_row}'] = "۵. مواد و کالاهای ورودی"
    ws[f'A{current_row}'].font = section_font
    ws[f'A{current_row}'].fill = section_fill
    ws[f'A{current_row}'].alignment = center_align
    current_row += 1
    
    # Materials headers
    mat_headers = ["نوع مصالح", "مقدار ورودی", "مقدار مصرفی", "محل نگهداری", "شماره بارنامه"]
    for col, header in enumerate(mat_headers, 1):
        cell = ws.cell(row=current_row, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    current_row += 1
    
    # Materials data
    for mat in form.incoming_materials:
        ws.cell(row=current_row, column=1, value=mat.material_type).border = thin_border
        ws.cell(row=current_row, column=1).alignment = right_align
        ws.cell(row=current_row, column=2, value=mat.incoming_amount).border = thin_border
        ws.cell(row=current_row, column=2).alignment = center_align
        ws.cell(row=current_row, column=3, value=mat.used_amount).border = thin_border
        ws.cell(row=current_row, column=3).alignment = center_align
        ws.cell(row=current_row, column=4, value=mat.storage_place or '').border = thin_border
        ws.cell(row=current_row, column=4).alignment = center_align
        ws.cell(row=current_row, column=5, value=mat.waybill_number or '').border = thin_border
        ws.cell(row=current_row, column=5).alignment = center_align
        current_row += 1
    
    current_row += 1
    
    # ==================== Section 6: Climate Condition ====================
    ws.merge_cells(f'A{current_row}:F{current_row}')
    ws[f'A{current_row}'] = "۶. شرایط آب و هوایی"
    ws[f'A{current_row}'].font = section_font
    ws[f'A{current_row}'].fill = section_fill
    ws[f'A{current_row}'].alignment = center_align
    current_row += 1
    
    # Climate headers
    climate_headers = ["حداقل دما", "حداکثر دما", "رطوبت (%)", "وضعیت آب و هوا", "سرعت باد", "تاثیر بر پروژه"]
    for col, header in enumerate(climate_headers, 1):
        cell = ws.cell(row=current_row, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    current_row += 1
    
    # Climate data
    if form.climate_conditions:
        climate = form.climate_conditions[0]
        ws.cell(row=current_row, column=1, value=climate.min_temperature).border = thin_border
        ws.cell(row=current_row, column=1).alignment = center_align
        ws.cell(row=current_row, column=2, value=climate.max_temperature).border = thin_border
        ws.cell(row=current_row, column=2).alignment = center_align
        ws.cell(row=current_row, column=3, value=climate.humidity).border = thin_border
        ws.cell(row=current_row, column=3).alignment = center_align
        ws.cell(row=current_row, column=4, value=weather_to_persian(climate.weather_type)).border = thin_border
        ws.cell(row=current_row, column=4).alignment = center_align
        ws.cell(row=current_row, column=5, value=wind_to_persian(climate.wind_speed)).border = thin_border
        ws.cell(row=current_row, column=5).alignment = center_align
        ws.cell(row=current_row, column=6, value=effect_to_persian(climate.climate_effect)).border = thin_border
        ws.cell(row=current_row, column=6).alignment = center_align
        current_row += 1
    else:
        for col in range(1, 7):
            ws.cell(row=current_row, column=col, value='').border = thin_border
            ws.cell(row=current_row, column=col).alignment = center_align
        current_row += 1
    
    current_row += 1
    
    # ==================== Section 7: Project Issues ====================
    ws.merge_cells(f'A{current_row}:D{current_row}')
    ws[f'A{current_row}'] = "۷. مسائل و مشکلات پروژه"
    ws[f'A{current_row}'].font = section_font
    ws[f'A{current_row}'].fill = section_fill
    ws[f'A{current_row}'].alignment = center_align
    current_row += 1
    
    # Issues headers
    issue_headers = ["نوع مشکل", "تاثیر", "محل (کیلومتر)", "توضیحات"]
    for col, header in enumerate(issue_headers, 1):
        cell = ws.cell(row=current_row, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    current_row += 1
    
    # Issues data
    for issue in form.project_issues:
        ws.cell(row=current_row, column=1, value=issue.issue_type).border = thin_border
        ws.cell(row=current_row, column=1).alignment = right_align
        ws.cell(row=current_row, column=2, value=issue.effect or '').border = thin_border
        ws.cell(row=current_row, column=2).alignment = center_align
        ws.cell(row=current_row, column=3, value=issue.location_station or '').border = thin_border
        ws.cell(row=current_row, column=3).alignment = center_align
        ws.cell(row=current_row, column=4, value=issue.notes or '').border = thin_border
        ws.cell(row=current_row, column=4).alignment = center_align
        current_row += 1
    
    current_row += 1
    
    # ==================== Section 8: Safety ====================
    ws.merge_cells(f'A{current_row}:D{current_row}')
    ws[f'A{current_row}'] = "۸. ایمنی"
    ws[f'A{current_row}'].font = section_font
    ws[f'A{current_row}'].fill = section_fill
    ws[f'A{current_row}'].alignment = center_align
    current_row += 1
    
    # Safety headers
    safety_headers = ["وضعیت ایمنی", "بازرسی ایمنی", "حادثه رخ داده", "توضیحات حادثه"]
    for col, header in enumerate(safety_headers, 1):
        cell = ws.cell(row=current_row, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    current_row += 1
    
    # Safety data
    if form.safety_records:
        safety = form.safety_records[0]
        ws.cell(row=current_row, column=1, value=safety_to_persian(safety.safety_situation)).border = thin_border
        ws.cell(row=current_row, column=1).alignment = center_align
        ws.cell(row=current_row, column=2, value='بله' if safety.safety_inspection else 'خیر').border = thin_border
        ws.cell(row=current_row, column=2).alignment = center_align
        ws.cell(row=current_row, column=3, value='بله' if safety.incident_occurred else 'خیر').border = thin_border
        ws.cell(row=current_row, column=3).alignment = center_align
        ws.cell(row=current_row, column=4, value=safety.incident_explanation or '').border = thin_border
        ws.cell(row=current_row, column=4).alignment = center_align
        current_row += 1
    else:
        for col in range(1, 5):
            ws.cell(row=current_row, column=col, value='').border = thin_border
            ws.cell(row=current_row, column=col).alignment = center_align
        current_row += 1
    
    current_row += 1
    
    # ==================== Section 9: Events ====================
    ws.merge_cells(f'A{current_row}:B{current_row}')
    ws[f'A{current_row}'] = "۹. رویدادها"
    ws[f'A{current_row}'].font = section_font
    ws[f'A{current_row}'].fill = section_fill
    ws[f'A{current_row}'].alignment = center_align
    current_row += 1
    
    # Events headers
    event_headers = ["نام رویداد", "توضیحات"]
    for col, header in enumerate(event_headers, 1):
        cell = ws.cell(row=current_row, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    ws.column_dimensions['B'].width = 50
    current_row += 1
    
    # Events data
    for event in form.events:
        ws.cell(row=current_row, column=1, value=event.event_name or '').border = thin_border
        ws.cell(row=current_row, column=1).alignment = center_align
        ws.cell(row=current_row, column=2, value=event.explanation or '').border = thin_border
        ws.cell(row=current_row, column=2).alignment = center_align
        current_row += 1
    
    # Adjust column widths
    ws.column_dimensions['A'].width = 25
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 20
    ws.column_dimensions['D'].width = 20
    ws.column_dimensions['E'].width = 25
    ws.column_dimensions['F'].width = 20
    ws.column_dimensions['G'].width = 20
    
    return wb
