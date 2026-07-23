"""
PPK Processing Result Processor
Processes PPK (Post-Processed Kinematic) GPS data files.

Updated with improved filtering logic using standard deviation and time-based filtering.
Added Excel export with colored groups feature.
"""

import pandas as pd
import numpy as np
from datetime import datetime
import os
from typing import Tuple, List, Dict, Optional
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils.dataframe import dataframe_to_rows
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


class PPKProcessor:
    """
    Process PPK (Post-Processed Kinematic) GPS result files.
    Uses improved filtering with standard deviation or minimum elevation threshold,
    and time-based filtering at both start and end of groups.
    """
    
    CODE_NAME = "PPK"
    NEEDED_COLS = ["Note", "Latitude", "Longitude", "Ell.Height (m)"]
    
    def __init__(self, file_path: str, time_threshold: float = 5.0, elevation_threshold: Optional[float] = None):
        """
        Initialize the PPK processor.
        
        Args:
            file_path: Path to the input CSV file
            time_threshold: Time threshold in minutes for filtering sparse data
            elevation_threshold: Optional minimum elevation threshold. If None, uses std deviation method
        """
        self.file_path = Path(file_path)
        self.time_threshold = time_threshold
        self.elevation_threshold = elevation_threshold
        
        # Load and process data
        self.data = self._load_data()
        self._extract_date()
        self.group_data = self._separate_data()
        
        # Store original data BEFORE filtering for Excel export
        self.original_data = self.data.copy()
        
        # Apply filtering
        if elevation_threshold is not None:
            self.removed_notes = self._remove_by_min_elevation(elevation_threshold)
        else:
            self.removed_notes = self._remove_by_std_deviation()
        
        self._remove_by_time(time_threshold=time_threshold)
        self._remove_empty_groups()
    
    def _load_data(self) -> pd.DataFrame:
        """
        Load CSV file and prepare data.
        
        Returns:
            Loaded DataFrame
        """
        if not self.file_path.exists():
            raise FileNotFoundError(f"The file {self.file_path} does not exist.")
        
        try:
            data = pd.read_csv(self.file_path)
            # Replace " " with np.nan for all columns
            data.replace(" ", pd.NA, inplace=True)
            # Convert columns to float, ignoring errors (NaNs will remain)
            data["Latitude"] = pd.to_numeric(data["Latitude"], errors="coerce")
            data["Longitude"] = pd.to_numeric(data["Longitude"], errors="coerce")
            data["Ell.Height (m)"] = pd.to_numeric(data["Ell.Height (m)"], errors="coerce")
            return data
        except Exception as e:
            raise ValueError(f"Error loading data: {e}")
    
    def _extract_date(self) -> None:
        """Parse datetime from Name column and add to DataFrame."""
        self.data[["id", "date", "time", "time_part"]] = self.data["Name"].str.split(" ", expand=True)
        self.data["datetime"] = pd.to_datetime(
            self.data["date"] + " " + self.data["time"] + " " + self.data["time_part"],
            format="%Y-%m-%d %I:%M:%S.%f %p"
        )
        self.data.sort_values(by=["datetime"], inplace=True, ignore_index=True)
    
    def _separate_data(self) -> List[pd.DataFrame]:
        """
        Separate data into groups based on Note column (where Note == 1 starts a new group).
        
        Returns:
            List of DataFrames, each representing a group
        """
        df = self.data.copy()
        
        start_indices = df.index[df['Note'] == 1].tolist()
        group_data = []
        
        for i, start in enumerate(start_indices):
            end = start_indices[i+1] if i+1 < len(start_indices) else len(df)
            group = df.iloc[start:end]
            group_data.append(group)
        
        return group_data
    
    def _remove_by_std_deviation(self) -> Dict:
        """
        Remove outlier points at start and end of each group using standard deviation.
        Points outside 2 standard deviations from mean elevation are removed.
        
        Returns:
            Dictionary of removed notes with their details
        """
        removed_notes = {}
        
        for idx, group in enumerate(self.group_data):
            elevs = group["Ell.Height (m)"].dropna().astype(float)
            if elevs.empty:
                removed_notes[idx+1] = []
                continue
            
            mean_elev = elevs.mean()
            std_elev = elevs.std()
            
            # Find first index close to mean (within 2 std)
            start = 0
            while start < len(group):
                elev = group.iloc[start]["Ell.Height (m)"]
                if pd.isna(elev):
                    start += 1
                    continue
                if abs(elev - mean_elev) <= 2 * std_elev:
                    break
                start += 1
            
            # Find last index close to mean (within 2 std)
            end = len(group) - 1
            while end >= start:
                elev = group.iloc[end]["Ell.Height (m)"]
                if pd.isna(elev):
                    end -= 1
                    continue
                if abs(elev - mean_elev) <= 2 * std_elev:
                    break
                end -= 1
            
            # Store removed notes
            removed = []
            for i in range(0, start):
                if pd.notna(group.iloc[i]["Ell.Height (m)"]):
                    removed.append({
                        "group_num": idx + 1,
                        "Note": int(group.iloc[i]["Note"]),
                        "Elevation": float(group.iloc[i]["Ell.Height (m)"])
                    })
            for i in range(end+1, len(group)):
                if pd.notna(group.iloc[i]["Ell.Height (m)"]):
                    removed.append({
                        "group_num": idx + 1,
                        "Note": int(group.iloc[i]["Note"]),
                        "Elevation": float(group.iloc[i]["Ell.Height (m)"])
                    })
            
            removed_notes[idx + 1] = removed
            group = group.iloc[start:end+1].reset_index(drop=True)
            self.group_data[idx] = group
        
        return removed_notes
    
    def _remove_by_min_elevation(self, min_elevation: float) -> Dict:
        """
        Remove points at start and end of each group below minimum elevation threshold.
        
        Args:
            min_elevation: Minimum elevation threshold
            
        Returns:
            Dictionary of removed notes with their details
        """
        removed_notes = {}
        
        for idx, group in enumerate(self.group_data):
            start = 0
            while start < len(group):
                elev = group.iloc[start]["Ell.Height (m)"]
                if pd.notna(elev) and elev < min_elevation:
                    start += 1
                else:
                    break
            
            end = len(group) - 1
            while end >= start:
                elev = group.iloc[end]["Ell.Height (m)"]
                if pd.notna(elev) and elev < min_elevation:
                    end -= 1
                else:
                    break
            
            # Store removed notes
            removed = []
            for i in range(0, start):
                if pd.notna(group.iloc[i]["Ell.Height (m)"]):
                    removed.append({
                        "group_num": idx + 1,
                        "Note": int(group.iloc[i]["Note"]),
                        "Elevation": float(group.iloc[i]["Ell.Height (m)"])
                    })
            for i in range(end+1, len(group)):
                if pd.notna(group.iloc[i]["Ell.Height (m)"]):
                    removed.append({
                        "group_num": idx + 1,
                        "Note": int(group.iloc[i]["Note"]),
                        "Elevation": float(group.iloc[i]["Ell.Height (m)"])
                    })
            
            removed_notes[idx + 1] = removed
            group = group.iloc[start:end+1].reset_index(drop=True)
            self.group_data[idx] = group
        
        return removed_notes
    
    def _remove_by_time(self, time_threshold: float) -> None:
        """
        Remove sparse data points at start and end based on time gaps.
        
        Args:
            time_threshold: Time threshold in minutes
        """
        for idx, group in enumerate(self.group_data):
            removed = []
            
            # Remove sparse data at the start
            start = 0
            while start + 1 < len(group):
                time_diff = (group.iloc[start + 1]["datetime"] - group.iloc[start]["datetime"]).total_seconds() / 60
                if time_diff > time_threshold:
                    # Remove all data up to and including start
                    for i in range(0, start + 1):
                        if pd.notna(group.iloc[i]["Ell.Height (m)"]):
                            removed.append({
                                "group_num": idx + 1,
                                "Note": int(group.iloc[i]["Note"]),
                                "Elevation": float(group.iloc[i]["Ell.Height (m)"])
                            })
                    group = group.iloc[start + 1:].reset_index(drop=True)
                    start = 0
                else:
                    start += 1
            
            # Remove sparse data at the end
            end = len(group) - 1
            while end > 0:
                time_diff = (group.iloc[end]["datetime"] - group.iloc[end - 1]["datetime"]).total_seconds() / 60
                if time_diff > time_threshold:
                    # Remove all data from end to the last
                    for i in range(end, len(group)):
                        if pd.notna(group.iloc[i]["Ell.Height (m)"]):
                            removed.append({
                                "group_num": idx + 1,
                                "Note": int(group.iloc[i]["Note"]),
                                "Elevation": float(group.iloc[i]["Ell.Height (m)"])
                            })
                    group = group.iloc[:end].reset_index(drop=True)
                    end = len(group) - 1
                else:
                    end -= 1
            
            # Add time-removed notes to existing removed_notes
            if removed:
                if idx + 1 in self.removed_notes:
                    self.removed_notes[idx + 1].extend(removed)
                else:
                    self.removed_notes[idx + 1] = removed
            
            if not group.empty:
                self.group_data[idx] = group
    
    def _remove_empty_groups(self) -> None:
        """Remove groups that are empty after filtering."""
        non_empty_groups = []
        for group in self.group_data:
            if not group.empty:
                non_empty_groups.append(group)
        self.group_data = non_empty_groups
    
    def _get_group_colors(self) -> List[str]:
        """
        Get a list of distinct colors for group coloring in Excel.
        
        Returns:
            List of hex color codes (without #)
        """
        # Distinct, visually appealing colors for groups
        colors = [
            "FFE6E6",  # Light Red
            "E6FFE6",  # Light Green
            "E6E6FF",  # Light Blue
            "FFFFE6",  # Light Yellow
            "FFE6FF",  # Light Magenta
            "E6FFFF",  # Light Cyan
            "FFD9B3",  # Light Orange
            "D9B3FF",  # Light Purple
            "B3FFD9",  # Light Mint
            "FFB3D9",  # Light Pink
            "D9FFB3",  # Light Lime
            "B3D9FF",  # Light Sky
            "FFDAB3",  # Peach
            "B3FFDA",  # Aquamarine
            "DAB3FF",  # Lavender
            "B3DAFF",  # Baby Blue
        ]
        return colors
    
    def generate_sorted_excel(self, output_dir: str) -> str:
        """
        Generate an Excel file with data sorted by datetime and Note,
        with each group (Note starting from 1) colored differently.
        Includes all original columns and splits datetime into separate columns.
        
        Args:
            output_dir: Directory to save the Excel file
            
        Returns:
            Path to the generated Excel file
        """
        if not OPENPYXL_AVAILABLE:
            raise ImportError("openpyxl is required for Excel export. Install it with: pip install openpyxl")
        
        os.makedirs(output_dir, exist_ok=True)
        
        base_name = os.path.basename(self.file_path)
        base_name_without_ext = os.path.splitext(base_name)[0]
        
        # Use ORIGINAL data (before filtering) for Excel export
        # Sort by datetime first, then by Note
        sorted_data = self.original_data.copy()
        sorted_data = sorted_data.sort_values(by=["datetime", "Note"]).reset_index(drop=True)
        
        # Extract datetime components (24-hour format)
        sorted_data["Year"] = sorted_data["datetime"].dt.year
        sorted_data["Month"] = sorted_data["datetime"].dt.month
        sorted_data["Day"] = sorted_data["datetime"].dt.day
        sorted_data["Hour"] = sorted_data["datetime"].dt.hour  # 24-hour format
        sorted_data["Minute"] = sorted_data["datetime"].dt.minute
        # Second with full decimal precision (seconds + microseconds)
        sorted_data["Second"] = sorted_data["datetime"].dt.second + sorted_data["datetime"].dt.microsecond / 1_000_000
        
        # Identify groups where Note resets to 1
        group_ids = []
        current_group = 0
        prev_note = None
        
        for idx, row in sorted_data.iterrows():
            note = row["Note"]
            if note == 1 and prev_note is not None:
                current_group += 1
            group_ids.append(current_group)
            prev_note = note
        
        sorted_data["Part"] = [g + 1 for g in group_ids]  # 1-based
        
        # Create Excel workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "PPK Sorted Data (Original)"
        
        # Get all original columns (excluding internal ones we added)
        internal_cols = ["datetime", "id", "date", "time", "time_part"]
        original_cols = [col for col in self.original_data.columns if col not in internal_cols]
        
        # Define columns to export: Date/Time components first, then all original columns, then Part
        date_time_cols = ["Year", "Month", "Day", "Hour", "Minute", "Second"]
        export_cols = date_time_cols + original_cols + ["Part"]
        
        # Create display names for headers
        display_names = {
            "Year": "Year",
            "Month": "Month", 
            "Day": "Day",
            "Hour": "Hour (24h)",
            "Minute": "Minute",
            "Second": "Second",
            "Note": "Note",
            "Latitude": "Latitude",
            "Longitude": "Longitude",
            "Ell.Height (m)": "Elevation (m)",
            "Part": "Part"
        }
        
        # Write header row with styling
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        header_alignment = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        for col_idx, col_name in enumerate(export_cols, 1):
            display_name = display_names.get(col_name, col_name)
            cell = ws.cell(row=1, column=col_idx, value=display_name)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # Get group colors
        colors = self._get_group_colors()
        
        # Write data rows with group coloring
        for row_idx, (_, row) in enumerate(sorted_data.iterrows(), 2):
            part = int(row["Part"])
            color = colors[(part - 1) % len(colors)]
            row_fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
            
            for col_idx, col in enumerate(export_cols, 1):
                value = row[col] if col in row.index else ""
                
                # Handle NaN values
                if pd.isna(value):
                    value = ""
                # Convert numeric types for Excel
                elif isinstance(value, (np.integer, np.floating)):
                    value = float(value) if isinstance(value, np.floating) else int(value)
                
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.fill = row_fill
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Adjust column widths based on content
        for col_idx, col_name in enumerate(export_cols, 1):
            col_letter = ws.cell(row=1, column=col_idx).column_letter
            # Set width based on column type
            if col_name in ["Year"]:
                ws.column_dimensions[col_letter].width = 8
            elif col_name in ["Month", "Day", "Hour", "Minute", "Second", "Note", "Part"]:
                ws.column_dimensions[col_letter].width = 10
            elif col_name in ["Latitude", "Longitude"]:
                ws.column_dimensions[col_letter].width = 15
            elif col_name in ["Ell.Height (m)"]:
                ws.column_dimensions[col_letter].width = 14
            else:
                ws.column_dimensions[col_letter].width = 15
        
        # Freeze the header row
        ws.freeze_panes = "A2"
        
        # Add a summary sheet
        ws_summary = wb.create_sheet(title="Summary")
        ws_summary["A1"] = "PPK Processing Summary"
        ws_summary["A1"].font = Font(bold=True, size=14)
        
        summary_data = [
            ["Total Points", len(sorted_data)],
            ["Number of Parts", len(sorted_data["Part"].unique())],
            ["Time Threshold (min)", self.time_threshold],
            ["Elevation Filter", self.elevation_threshold if self.elevation_threshold else "Standard Deviation"],
            ["", ""],
            ["Columns Included", len(export_cols)],
        ]
        
        for row_idx, (label, value) in enumerate(summary_data, 3):
            ws_summary.cell(row=row_idx, column=1, value=label).font = Font(bold=True)
            ws_summary.cell(row=row_idx, column=2, value=value)
        
        ws_summary.column_dimensions["A"].width = 25
        ws_summary.column_dimensions["B"].width = 20
        ws_summary.column_dimensions["C"].width = 15
        
        # Check if Solution_type column exists (needed for multiple sections)
        solution_col = None
        possible_cols = ['Solution_type', 'solution_type', 'SolutionType', 'Solution Type', 
                         'Type', 'Fix_Type', 'FixType', 'Quality', 'Status']
        for col in possible_cols:
            if col in sorted_data.columns:
                solution_col = col
                break
        
        # Get unique solution types for headers (including empty/NaN as "Empty")
        solution_types = []
        empty_count_total = 0
        if solution_col and solution_col in sorted_data.columns:
            # Get non-null unique values
            solution_types = sorted([str(x) for x in sorted_data[solution_col].dropna().unique()])
            # Count empty/NaN values
            empty_count_total = sorted_data[solution_col].isna().sum()
            # Add "Empty" to the list if there are any empty values
            if empty_count_total > 0:
                solution_types.append("Empty")
        
        # Add Part Statistics section with Solution Type breakdown
        current_row = 11
        ws_summary.cell(row=current_row, column=1, value="Part Statistics").font = Font(bold=True, size=12)
        current_row += 1
        
        unique_parts = sorted(sorted_data["Part"].unique())
        part_counts = sorted_data.groupby("Part").size()
        
        # Header for part stats
        ws_summary.cell(row=current_row, column=1, value="Part").font = Font(bold=True)
        ws_summary.cell(row=current_row, column=1).border = thin_border
        ws_summary.cell(row=current_row, column=2, value="Total Points").font = Font(bold=True)
        ws_summary.cell(row=current_row, column=2).border = thin_border
        
        # Add solution type columns in header
        if solution_types:
            for col_idx, sol_type in enumerate(solution_types, 3):
                cell = ws_summary.cell(row=current_row, column=col_idx, value=str(sol_type))
                cell.font = Font(bold=True)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center")
                ws_summary.column_dimensions[cell.column_letter].width = 12
        current_row += 1
        
        # Data rows for each part
        for part in unique_parts:
            color = colors[(part - 1) % len(colors)]
            
            # Part name with color
            cell_part = ws_summary.cell(row=current_row, column=1, value=f"Part {part}")
            cell_part.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
            cell_part.border = thin_border
            
            # Total count
            cell_count = ws_summary.cell(row=current_row, column=2, value=int(part_counts[part]))
            cell_count.border = thin_border
            cell_count.alignment = Alignment(horizontal="center")
            
            # Solution type counts per part
            if solution_col and solution_types:
                part_data = sorted_data[sorted_data["Part"] == part]
                part_solution_counts = part_data[solution_col].value_counts(dropna=True)
                part_empty_count = part_data[solution_col].isna().sum()
                
                for col_idx, sol_type in enumerate(solution_types, 3):
                    if sol_type == "Empty":
                        count = int(part_empty_count)
                    else:
                        count = int(part_solution_counts.get(sol_type, 0))
                    cell = ws_summary.cell(row=current_row, column=col_idx, value=count)
                    cell.border = thin_border
                    cell.alignment = Alignment(horizontal="center")
            
            current_row += 1
        
        # Add Total row
        cell_total_label = ws_summary.cell(row=current_row, column=1, value="TOTAL")
        cell_total_label.font = Font(bold=True)
        cell_total_label.border = thin_border
        
        cell_total_count = ws_summary.cell(row=current_row, column=2, value=len(sorted_data))
        cell_total_count.font = Font(bold=True)
        cell_total_count.border = thin_border
        cell_total_count.alignment = Alignment(horizontal="center")
        
        if solution_col and solution_types:
            total_solution_counts = sorted_data[solution_col].value_counts(dropna=True)
            for col_idx, sol_type in enumerate(solution_types, 3):
                if sol_type == "Empty":
                    count = int(empty_count_total)
                else:
                    count = int(total_solution_counts.get(sol_type, 0))
                cell = ws_summary.cell(row=current_row, column=col_idx, value=count)
                cell.font = Font(bold=True)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center")
        
        current_row += 2
        
        # Add Overall Solution Type Statistics section
        if solution_col and solution_col in sorted_data.columns:
            ws_summary.cell(row=current_row, column=1, value="Overall Solution Type Summary").font = Font(bold=True, size=12)
            current_row += 1
            
            # Header for solution type stats
            ws_summary.cell(row=current_row, column=1, value="Solution Type").font = Font(bold=True)
            ws_summary.cell(row=current_row, column=2, value="Count").font = Font(bold=True)
            ws_summary.cell(row=current_row, column=3, value="Percentage").font = Font(bold=True)
            ws_summary.cell(row=current_row, column=1).border = thin_border
            ws_summary.cell(row=current_row, column=2).border = thin_border
            ws_summary.cell(row=current_row, column=3).border = thin_border
            current_row += 1
            
            # Count each solution type with percentage (including empty/NaN)
            solution_counts = sorted_data[solution_col].value_counts(dropna=True)
            total_points = len(sorted_data)
            
            # First show non-empty solution types
            for solution_type, count in solution_counts.items():
                cell_type = ws_summary.cell(row=current_row, column=1, value=str(solution_type))
                cell_type.border = thin_border
                
                cell_count = ws_summary.cell(row=current_row, column=2, value=int(count))
                cell_count.border = thin_border
                cell_count.alignment = Alignment(horizontal="center")
                
                percentage = (count / total_points) * 100
                cell_pct = ws_summary.cell(row=current_row, column=3, value=f"{percentage:.1f}%")
                cell_pct.border = thin_border
                cell_pct.alignment = Alignment(horizontal="center")
                current_row += 1
            
            # Add empty/NaN count row if there are any
            if empty_count_total > 0:
                cell_type = ws_summary.cell(row=current_row, column=1, value="Empty")
                cell_type.border = thin_border
                cell_type.font = Font(italic=True)
                
                cell_count = ws_summary.cell(row=current_row, column=2, value=int(empty_count_total))
                cell_count.border = thin_border
                cell_count.alignment = Alignment(horizontal="center")
                
                percentage = (empty_count_total / total_points) * 100
                cell_pct = ws_summary.cell(row=current_row, column=3, value=f"{percentage:.1f}%")
                cell_pct.border = thin_border
                cell_pct.alignment = Alignment(horizontal="center")
                current_row += 1
        else:
            ws_summary.cell(row=current_row, column=1, value="No Solution Type column found in data")
            ws_summary.cell(row=current_row, column=1).font = Font(italic=True, color="888888")
        
        # Save the workbook
        excel_file = os.path.join(output_dir, f"{self.CODE_NAME}_{base_name_without_ext}_sorted_colored.xlsx")
        wb.save(excel_file)
        
        return excel_file
    
    def get_sorted_data_json(self) -> List[Dict]:
        """
        Get sorted data as JSON for web display.
        Uses ORIGINAL data (before filtering) for display.
        
        Returns:
            List of dictionaries with sorted data and group information
        """
        # Use ORIGINAL data (before filtering) for web display
        sorted_data = self.original_data.sort_values(by=["datetime", "Note"]).reset_index(drop=True)
        
        # Identify groups
        group_ids = []
        current_group = 0
        prev_note = None
        
        for idx, row in sorted_data.iterrows():
            note = row["Note"]
            if note == 1 and prev_note is not None:
                current_group += 1
            group_ids.append(current_group)
            prev_note = note
        
        sorted_data["group_id"] = group_ids
        
        # Convert to JSON-serializable format
        result = []
        for _, row in sorted_data.iterrows():
            result.append({
                "datetime": row["datetime"].strftime("%Y-%m-%d %I:%M:%S %p") if pd.notna(row["datetime"]) else "",
                "note": int(row["Note"]) if pd.notna(row["Note"]) else None,
                "latitude": float(row["Latitude"]) if pd.notna(row["Latitude"]) else None,
                "longitude": float(row["Longitude"]) if pd.notna(row["Longitude"]) else None,
                "elevation": float(row["Ell.Height (m)"]) if pd.notna(row["Ell.Height (m)"]) else None,
                "group_id": int(row["group_id"]) + 1  # 1-based for display
            })
        
        return result
    
    def save_files(self, output_dir: str, start_num: int = 1) -> Tuple[str, str, List[str], str]:
        """
        Save processed PPK data to files.
        
        Args:
            output_dir: Directory to save output files
            start_num: Starting number for output file naming
            
        Returns:
            Tuple of (ppk_file, removed_file, text_files, excel_file) paths
        """
        os.makedirs(output_dir, exist_ok=True)
        
        base_name = os.path.basename(self.file_path)
        base_name_without_ext = os.path.splitext(base_name)[0]
        
        # Save main CSV file with all data
        ppk_file = os.path.join(output_dir, f"{self.CODE_NAME}_{base_name_without_ext}_processed.csv")
        all_data = pd.concat(self.group_data, ignore_index=True)
        all_data[self.NEEDED_COLS].dropna().to_csv(ppk_file, index=False, header=True)
        
        # Save individual group text files
        text_files = []
        for idx, group in enumerate(self.group_data):
            rows = group[self.NEEDED_COLS].dropna().shape[0]
            if rows > 0:
                out_file = os.path.join(output_dir, f"{start_num + idx}_{rows}.txt")
                text_files.append(out_file)
                # Save as .txt without header or index
                group[self.NEEDED_COLS].dropna().to_csv(out_file, index=False, header=False)
        
        # Save removed notes file
        removed_file = os.path.join(output_dir, f"{self.CODE_NAME}_{base_name_without_ext}_removed_notes.csv")
        removed_rows = []
        for group_num, notes in self.removed_notes.items():
            for note in notes:
                removed_rows.append({
                    "Group": note["group_num"],
                    "Note": note["Note"],
                    "Elevation": note["Elevation"]
                })
        
        if removed_rows:
            removed_data = pd.DataFrame(removed_rows)
            removed_data.to_csv(removed_file, index=False, header=True)
        else:
            # Create empty file with headers
            pd.DataFrame(columns=["Group", "Note", "Elevation"]).to_csv(removed_file, index=False, header=True)
        
        # Generate the sorted colored Excel file
        excel_file = ""
        if OPENPYXL_AVAILABLE:
            try:
                excel_file = self.generate_sorted_excel(output_dir)
            except Exception as e:
                print(f"Warning: Could not generate Excel file: {e}")
                excel_file = ""
        else:
            print("Warning: openpyxl not available, skipping Excel export")
        
        print(f"PPK processing complete. Saved {len(text_files)} group files.")
        print(f"Removed notes: {self.removed_notes}")
        
        return ppk_file, removed_file, text_files, excel_file
    
    def get_results(self) -> dict:
        """
        Get processing results summary.
        
        Returns:
            Dictionary with processing statistics
        """
        total_removed = sum(len(notes) for notes in self.removed_notes.values())
        total_points = sum(len(group) for group in self.group_data)
        
        return {
            'total_points': total_points,
            'groups': len(self.group_data),
            'removed_points': total_removed,
            'time_threshold': self.time_threshold,
            'elevation_threshold': self.elevation_threshold if self.elevation_threshold else 'std_deviation',
            'removed_notes': self.removed_notes
        }


# Backward compatibility alias
PPK_Processing_Result_DataLoader = PPKProcessor
